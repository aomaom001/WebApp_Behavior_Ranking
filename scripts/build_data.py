#!/usr/bin/env python3
"""
build_data.py — generate frontend/data/data.json (+ frontend/data/detail/*.json) from
the MATELINE "TICKET CLOSED" Excel exports.

The aggregation itself lives in agg_core.py (shared with the PostgreSQL-backed API),
so the numbers are identical whatever the source. This file only reads the Excel sheets
and writes the JSON the static dashboard loads.

Usage:
    python scripts/build_data.py --src "C:/path/to/excel/folder" --out frontend/data/data.json
    python scripts/build_data.py --src .. --validate frontend/data/data.json   # dry-run + compare

Requires: openpyxl  (pip install openpyxl)

Input  : files named "..._YYYYMM.xlsx", sheet "RAW Data", one row per WO.
Output : { months, records[ {..., bd:{sevT,status,wtype,root,sla}} ], behaviors, dims, meta }
         plus data/detail/<month>.json (per-WO rows for drill-down).
"""

import argparse
import datetime
import glob
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import agg_core  # noqa: E402  (canonical aggregation logic)

SHEET = "RAW Data"


def month_from_filename(path):
    """'MATELINE TICKET CLOSED_202507.xlsx' -> '2025-07'."""
    m = re.search(r"(\d{4})(\d{2})", os.path.basename(path))
    if not m:
        raise ValueError(f"cannot find YYYYMM in filename: {path}")
    return f"{m.group(1)}-{m.group(2)}"


def process_file(path, month):
    """Read one monthly Excel export -> (records, detail_rows) via agg_core."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError(f"sheet '{SHEET}' not found in {path}; has {wb.sheetnames}")
    it = wb[SHEET].iter_rows(values_only=True)
    header = list(next(it))
    missing = [c for c in agg_core.NEEDED if c not in header]
    if missing:
        wb.close()
        raise ValueError(f"{path} missing columns: {missing}")
    rows = [dict(zip(header, r)) for r in it]
    wb.close()
    return agg_core.aggregate_month(rows, month, has_site=agg_core.OPT_SITE in header)


def write_detail(path, month, detail_rows):
    obj = agg_core.encode_detail(month, detail_rows)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))
    return os.path.getsize(path)


def build(src_dir, detail_dir=None):
    files = sorted(glob.glob(os.path.join(src_dir, "*.xlsx")))
    files = [f for f in files
             if not os.path.basename(f).startswith("~$")        # skip Excel lock files
             and re.search(r"\d{6}", os.path.basename(f))]
    if not files:
        sys.exit(f"no '*YYYYMM*.xlsx' files found in {src_dir}")
    all_records, months, sources, total_wo = [], [], [], 0
    for f in files:
        month = month_from_filename(f)
        print(f"  reading {os.path.basename(f)}  ->  {month}", flush=True)
        recs, detail = process_file(f, month)
        teams = len(recs)
        wo = sum(r["wo"] for r in recs)
        tickets = sum(r["tickets"] for r in recs)
        if detail_dir:
            kb = write_detail(os.path.join(detail_dir, month + ".json"), month, detail) / 1024
            print(f"    detail: {len(detail)} WO -> {month}.json ({kb:.0f} KB)", flush=True)
        st = os.stat(f)
        sources.append({
            "file": os.path.basename(f), "month": month, "teams": teams, "wo": wo, "tickets": tickets,
            "size_kb": round(st.st_size / 1024),
            "modified": datetime.datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d"),
        })
        total_wo += wo
        all_records.extend(recs)
        months.append(month)
        print(f"    {teams} teams, {wo} WO", flush=True)
    months = sorted(set(months))
    meta = {
        "generated_at": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "src": os.path.abspath(src_dir),
        "source_system": "ระบบ MATELINE — รายงาน Ticket Closed (.xlsx) รายเดือน",
        "total_files": len(files), "total_records": len(all_records), "total_wo": total_wo,
        "sources": sources,
    }
    return {"months": months, "records": all_records, "behaviors": agg_core.BEHAVIORS,
            "dims": agg_core.build_dims(all_records), "meta": meta}


def validate(generated, ref_path):
    ref = json.load(open(ref_path, encoding="utf-8"))
    ref_by = {(r["team"], r["m"]): r for r in ref["records"]}
    gen_by = {(r["team"], r["m"]): r for r in generated["records"]}
    common = set(ref_by) & set(gen_by)
    print(f"\nVALIDATION vs {ref_path}")
    print(f"  reference records: {len(ref['records'])} | generated: {len(generated['records'])} | common keys: {len(common)}")
    for fld in ["wo", "tickets", "dup", "nowork", "cancel", "sys", "man", "cross", "name", "region", "prov", "skill"]:
        ok = sum(1 for k in common if ref_by[k].get(fld) == gen_by[k].get(fld))
        pct = ok / len(common) * 100 if common else 0
        print(f"  {'OK ' if ok == len(common) else '~~ '}{fld:8} {ok}/{len(common)}  ({pct:.1f}%)")


def main():
    ap = argparse.ArgumentParser(description="Build data/data.json from MATELINE Excel exports.")
    ap.add_argument("--src", required=True, help="folder containing the *_YYYYMM.xlsx files")
    ap.add_argument("--out", help="output JSON path (omit when only --validate)")
    ap.add_argument("--validate", metavar="REF.json", help="compare result to an existing data.json")
    args = ap.parse_args()

    print(f"Building from: {os.path.abspath(args.src)}")
    detail_dir = os.path.join(os.path.dirname(args.out) or ".", "detail") if args.out else None
    data = build(args.src, detail_dir)
    print(f"\nTotal: {len(data['records'])} records across {len(data['months'])} months")
    if args.validate:
        validate(data, args.validate)
    if args.out:
        os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
        with open(args.out, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
        print(f"\nWrote {args.out}  ({os.path.getsize(args.out)/1024/1024:.2f} MB)")


if __name__ == "__main__":
    main()
