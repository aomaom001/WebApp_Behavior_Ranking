#!/usr/bin/env python3
"""
load_work_load.py — load the "Work Load Per Days" Excel exports into PostgreSQL (table work_load).

Usage:
    PG_HOST=localhost PG_PORT=5440 PG_DB=tt PG_USER=postgres PG_PASSWORD=ttpass \
        python scripts/load_work_load.py --src "C:/.../Dashboard/Work Load"

Requires: openpyxl, psycopg
"""

import argparse
import datetime
import glob
import os
import re
import sys

try:
    import openpyxl
    import psycopg
except ImportError as e:
    sys.exit(f"missing dependency: {e}.  pip install openpyxl \"psycopg[binary]\"")

# (Excel header, snake_case column, kind)   kind: ts=timestamp, num=numeric, txt=text
COLS = [
    ("Flag Month", "flag_month", "ts"), ("Flag WK", "flag_wk", "txt"),
    ("Flag Date final", "flag_date_final", "ts"), ("Flag Date1", "flag_date1", "ts"),
    ("Status Team", "status_team", "txt"), ("Region", "region", "txt"),
    ("Date", "date", "ts"), ("Team", "team", "txt"), ("Skill", "skill", "txt"),
    ("Manager", "manager", "txt"), ("Province", "province", "txt"), ("Suspend", "suspend", "txt"),
    ("Team Count", "team_count", "num"), ("Dispatched", "dispatched", "num"),
    ("Accepted", "accepted", "num"), ("Departed", "departed", "num"), ("Arrived", "arrived", "num"),
    ("Completed", "completed", "num"), ("Closed (Leaved)", "closed_leaved", "num"),
    ("% Closed", "pct_closed", "num"), ("Work Days", "work_days", "num"),
    ("PDT / Day (AVG)", "pdt_day_avg", "num"), ("PDT / Day (MEAN)", "pdt_day_mean", "num"),
    ("Total Team PDT>= 2.5", "total_team_pdt_ge25", "num"), ("Total Team PDT < 2.5", "total_team_pdt_lt25", "num"),
    ("% SLA", "pct_sla", "num"), ("Man Hour", "man_hour", "num"), ("Man Hour / Day", "man_hour_day", "num"),
]
SQLTYPE = {"ts": "timestamp", "num": "double precision", "txt": "text"}


def conninfo():
    return (f"host={os.environ.get('PG_HOST', 'localhost')} port={os.environ.get('PG_PORT', '5432')} "
            f"dbname={os.environ.get('PG_DB', 'postgres')} user={os.environ.get('PG_USER', 'postgres')} "
            f"password={os.environ.get('PG_PASSWORD', '')}")


def coerce(kind, v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    if kind == "ts":
        return v if isinstance(v, (datetime.datetime, datetime.date)) else None
    if kind == "num":
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    return str(v)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True, help="folder with the Work Load *_YYYYMM.xlsx files")
    args = ap.parse_args()
    files = [f for f in sorted(glob.glob(os.path.join(args.src, "*.xlsx")))
             if not os.path.basename(f).startswith("~$") and re.search(r"\d{6}", os.path.basename(f))]
    if not files:
        sys.exit(f"no '*YYYYMM*.xlsx' files in {args.src}")

    snake = [c[1] for c in COLS] + ["source_month"]
    coldefs = ", ".join(f"{c[1]} {SQLTYPE[c[2]]}" for c in COLS) + ", source_month text"
    norm = lambda s: re.sub(r"\s+", " ", str(s)).strip().lower()
    want = {norm(h): (snk, kind) for h, snk, kind in COLS}

    print(f"Loading {len(files)} files into work_load ...", flush=True)
    with psycopg.connect(conninfo()) as conn, conn.cursor() as cur:
        cur.execute("DROP TABLE IF EXISTS work_load")
        cur.execute(f"CREATE TABLE work_load ({coldefs})")
        conn.commit()
        total = 0
        for f in files:
            m = re.search(r"(\d{4})(\d{2})", os.path.basename(f))
            src_month = f"{m.group(1)}-{m.group(2)}" if m else None
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            it = ws.iter_rows(values_only=True)
            header = list(next(it))
            # position of each wanted column in THIS sheet (matched on normalised header)
            pos = {}
            for i, h in enumerate(header):
                key = norm(h)
                if key in want:
                    pos[want[key][0]] = (i, want[key][1])
            missing = [c[1] for c in COLS if c[1] not in pos]
            if missing:
                print(f"  ! {os.path.basename(f)} missing columns: {missing}", flush=True)
            n = 0
            with cur.copy(f"COPY work_load ({', '.join(snake)}) FROM STDIN") as cp:
                for row in it:
                    out = []
                    for c in COLS:
                        if c[1] in pos:
                            i, kind = pos[c[1]]
                            out.append(coerce(kind, row[i] if i < len(row) else None))
                        else:
                            out.append(None)
                    out.append(src_month)
                    cp.write_row(out)
                    n += 1
            wb.close()
            conn.commit()
            total += n
            print(f"  {os.path.basename(f)}: {n} rows", flush=True)
        cur.execute("CREATE INDEX IF NOT EXISTS ix_wl_month ON work_load (source_month)")
        cur.execute("CREATE INDEX IF NOT EXISTS ix_wl_team ON work_load (team)")
        conn.commit()
    print(f"DONE — {total} rows in work_load", flush=True)


if __name__ == "__main__":
    main()
