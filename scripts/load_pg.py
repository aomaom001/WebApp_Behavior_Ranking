#!/usr/bin/env python3
"""
load_pg.py — load the MATELINE "TICKET CLOSED" Excel exports into PostgreSQL.

Creates table `mateline_ticket_closed` (snake_case columns) and bulk-COPYs every WO row.
Includes the columns the live API needs (for the duplicate-WO ranking) plus the timestamp
and point columns used by the PDT & Point detection (arrived/completed/point/created_time).

Usage:
    PG_HOST=localhost PG_PORT=5440 PG_DB=tt PG_USER=postgres PG_PASSWORD=... \
        python scripts/load_pg.py --src "C:/path/to/excel/folder"

Requires: openpyxl, psycopg  (pip install openpyxl "psycopg[binary]")
"""

import argparse
import datetime
import glob
import os
import re
import sys

try:
    import openpyxl
    import psycopg
except ImportError as e:
    sys.exit(f"missing dependency: {e}.  pip install openpyxl \"psycopg[binary]\"")

SHEET = "RAW Data"
# (Excel header, snake_case column, kind)  kind: ts=timestamp, num=numeric, txt=text
COLS = [
    ("Work Order ID", "work_order_id", "txt"), ("Team", "team", "txt"),
    ("Source Ticket ID", "source_ticket_id", "txt"), ("Province", "province", "txt"),
    ("Region", "region", "txt"), ("Skill", "skill", "txt"), ("Status", "status", "txt"),
    ("WO Creator", "wo_creator", "txt"), ("Complete Solution", "complete_solution", "txt"),
    ("Severity", "severity", "txt"), ("Work Type", "work_type", "txt"),
    ("Root Cause", "root_cause", "txt"), ("SLA", "sla", "txt"), ("Site ID", "site_id", "txt"),
    ("Created Time", "created_time", "ts"), ("Arrived", "arrived", "ts"),
    ("Completed", "completed", "ts"), ("Closed", "closed", "ts"), ("Point", "point", "num"),
]
SQLTYPE = {"ts": "timestamp", "num": "double precision", "txt": "text"}


def conninfo():
    return (f"host={os.environ.get('PG_HOST', 'localhost')} port={os.environ.get('PG_PORT', '5432')} "
            f"dbname={os.environ.get('PG_DB', 'postgres')} user={os.environ.get('PG_USER', 'postgres')} "
            f"password={os.environ.get('PG_PASSWORD', '')}")


def coerce(kind, v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    if kind == "ts":
        if isinstance(v, (datetime.datetime, datetime.date)):
            return v
        s = str(v).strip()
        for fmt in (None, "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                return datetime.datetime.fromisoformat(s) if fmt is None else datetime.datetime.strptime(s, fmt)
            except ValueError:
                continue
        return None
    if kind == "num":
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    return str(v)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True, help="folder with the *_YYYYMM.xlsx files")
    args = ap.parse_args()
    files = [f for f in sorted(glob.glob(os.path.join(args.src, "*.xlsx")))
             if not os.path.basename(f).startswith("~$") and re.search(r"\d{6}", os.path.basename(f))]
    if not files:
        sys.exit(f"no '*YYYYMM*.xlsx' files in {args.src}")

    snake = [c[1] for c in COLS]
    coldefs = ", ".join(f"{c[1]} {SQLTYPE[c[2]]}" for c in COLS)
    norm = lambda s: re.sub(r"\s+", " ", str(s)).strip().lower()
    want = {norm(h): (snk, kind) for h, snk, kind in COLS}
    print(f"Loading {len(files)} files into mateline_ticket_closed ...", flush=True)
    with psycopg.connect(conninfo()) as conn, conn.cursor() as cur:
        cur.execute("DROP TABLE IF EXISTS mateline_ticket_closed")
        cur.execute(f"CREATE TABLE mateline_ticket_closed ({coldefs})")
        conn.commit()
        total = 0
        for f in files:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            if SHEET not in wb.sheetnames:
                wb.close(); print(f"  SKIP {os.path.basename(f)} (no '{SHEET}' sheet)"); continue
            it = wb[SHEET].iter_rows(values_only=True)
            header = list(next(it))
            pos = {}
            for i, h in enumerate(header):
                k = norm(h)
                if k in want:
                    pos[want[k][0]] = (i, want[k][1])
            missing = [c[1] for c in COLS if c[1] not in pos]
            if missing:
                print(f"  ! {os.path.basename(f)} missing: {missing}", flush=True)
            n = 0
            with cur.copy(f"COPY mateline_ticket_closed ({', '.join(snake)}) FROM STDIN") as cp:
                for row in it:
                    out = []
                    for c in COLS:
                        if c[1] in pos:
                            i, kind = pos[c[1]]
                            out.append(coerce(kind, row[i] if i < len(row) else None))
                        else:
                            out.append(None)
                    cp.write_row(out)
                    n += 1
            wb.close()
            conn.commit()
            total += n
            print(f"  {os.path.basename(f)}: {n} rows", flush=True)
        cur.execute("CREATE INDEX IF NOT EXISTS ix_mtc_closed ON mateline_ticket_closed (closed)")
        cur.execute("CREATE INDEX IF NOT EXISTS ix_mtc_team ON mateline_ticket_closed (team)")
        conn.commit()
    print(f"DONE — {total} rows in mateline_ticket_closed", flush=True)


if __name__ == "__main__":
    main()
